# -*- coding: utf-8 -*-
"""生成《设备采购清单》（按指导老师报销范围调整）。
只保留可报销的电子测量器件；插座/插排由实验室提供；辅材不申请。
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT_CN = "微软雅黑"
FONT_EN = "Calibri"


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = FONT_EN
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)


def style_doc(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_EN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    for i in range(1, 4):
        st = doc.styles[f"Heading {i}"]
        st.font.name = FONT_EN
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)


def para(doc, text="", size=10.5, bold=False, align=None, space_after=3, indent_first=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if indent_first:
        p.paragraph_format.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    set_font(r, size, bold)
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_font(r, size)
    return p


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_table(doc, headers, rows, widths=None, header_fill="1F3A5F", total_rows=(), warn_rows=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        table._tbl.tblPr.append(layout)
        for i, w in enumerate(widths):
            table.columns[i].width = Cm(w)

    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_font(r, 9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(hdr[i], header_fill)

    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, val in enumerate(row):
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci in (0, 3, 4, 5) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val))
            set_font(r, 8.5)
        if ri in total_rows:
            for c in cells:
                shade_cell(c, "EAF0F6")
                for r in c.paragraphs[0].runs:
                    r.font.bold = True
        if ri in warn_rows:
            for c in cells:
                shade_cell(c, "FFF6E5")
    return table


def build(path):
    doc = Document()
    style_doc(doc)
    sec = doc.sections[0]
    sec.top_margin = Cm(1.9)
    sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    para(doc, "设备采购清单", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, "项目：EdgeGuard · 第十一届华为 ICT 大赛创新赛 赛题1",
         size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, "按指导老师 2026-10-23 的报销范围意见调整",
         size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    # 一、调整说明
    doc.add_heading("一、本次调整说明", level=1)
    para(doc, "根据指导老师意见，本次只申请可报销的电子测量器件，其余物品不再采购：", indent_first=True, space_after=3)
    bullet(doc, "只申请两类：单相电能计量模块（电表）、USB 转 TTL 串口模块。")
    bullet(doc, "插座、插排：实验室已有，可直接使用，不再采购。")
    bullet(doc, "扎带、绝缘胶带、导线、端子、备用耗材：无法作为实验材料报销，不再申请。")
    bullet(doc, "万用表：如实验室可借用则不买；如无，再单独申请。")

    # 二、采购清单
    doc.add_heading("二、采购清单（可报销项）", level=1)
    headers = ["序号", "物品名称", "规格/型号", "数量", "单价\n(元)", "小计\n(元)", "在哪买", "报销"]
    rows = [
        ["1", "单相电能计量模块", "PZEM-004T v3.0\n（含开合式 CT 互感器）", "1", "53.00", "53.00",
         "淘宝/天猫搜索「PZEM-004T v3.0 开合式」，选销量高、可开票的电子元件店", "✓ 可报销\n（电表）"],
        ["2", "同上（备用）", "PZEM-004T v3.0", "1", "53.00", "53.00",
         "同上（建议与第 1 项同店购买，省运费）", "✓ 可报销\n（经费紧可只买 1 个）"],
        ["3", "USB 转 TTL 模块", "CH340X 或 CP2102", "2", "12.00", "24.00",
         "立创商城（CH340X 模块，商品编号 C49247784）；或淘宝同款", "✓ 可报销\n（USB 转模块）"],
        ["", "小计", "——", "——", "——", "130.00", "——", "——"],
        ["4", "数字万用表", "得力 DL8490（入门级）", "1", "37.15", "37.15",
         "嘉立创 FA（jlcfa.com，含税价）", "？待确认\n（实验室可借用则不买）"],
        ["", "合计（含万用表）", "——", "——", "——", "167.15", "——", "——"],
    ]
    add_table(doc, headers, rows,
              widths=[1.1, 2.6, 3.2, 0.9, 1.2, 1.3, 3.9, 2.4],
              total_rows=(3, 5), warn_rows=(4,))
    para(doc, "说明：价格为 2026 年 10 月 23 日核实的参考价（PZEM 模块取自阿里巴巴批发报价折算零售，"
              "万用表取自嘉立创含税价），最终以实际发票金额为准。",
         size=8.5, space_after=6)

    # 三、不采购（另起一页，避免表格孤行）
    doc.add_page_break()
    doc.add_heading("三、本次不采购的物品", level=1)
    skip_headers = ["物品", "不采购原因"]
    skip_rows = [
        ["插座盒、插排", "实验室已有，可直接使用（老师明确）"],
        ["尼龙扎带", "老师明确无法作为实验材料报销"],
        ["绝缘胶带、铜芯导线、接线端子", "属通用辅材，不在报销范围；实验室应可提供"],
        ["备用耗材（保险丝等）", "同上，不再申请"],
    ]
    add_table(doc, skip_headers, skip_rows, widths=[4.6, 12.0])

    # 四、在哪买与发票
    doc.add_heading("四、在哪买与发票要求", level=1)
    para(doc, "1. 平台选择：优先淘宝/天猫电子元件店、立创商城、嘉立创 FA 等可开票渠道；"
              "同一家能买齐的尽量同店下单，减少运费。", indent_first=True, space_after=2)
    para(doc, "2. 开票要求：下单时备注「需开增值税普通发票」，并保留订单与支付记录；"
              "电子发票请另存 PDF。", indent_first=True, space_after=2)
    para(doc, "3. 发票抬头与税号：请老师提供；若学校要求走指定采购平台，请告知，我们改在该平台下单。",
         indent_first=True, space_after=2)
    para(doc, "4. 收货后：保留包装、说明书与型号标签，作为比赛材料与资产登记凭证。",
         indent_first=True)

    # 五、用途与安全
    doc.add_heading("五、用途与安全说明", level=1)
    para(doc, "用途：电能计量模块用于采集电压、电流、功率、功率因数，为四类负载（电瓶车充电器、"
              "电热水壶、吹风机、正常待机）建立波形数据集，训练识别模型并产出初赛所需的量化指标。",
         indent_first=True, space_after=2)
    para(doc, "安全：模块仅在低压测量侧使用，接线全程断电操作、接后先验电；实验在实验室进行，"
              "不擅自改动实验室配电线路。", indent_first=True)

    # 六、审批
    doc.add_heading("六、审批", level=1)
    para(doc, "指导教师意见：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿", space_after=8)
    para(doc, "签字：＿＿＿＿＿＿＿        日期：＿＿＿＿年＿＿月＿＿日", space_after=8)
    para(doc, "申请人（学生签名）：＿＿＿＿＿＿＿  ＿＿＿＿＿＿＿  ＿＿＿＿＿＿＿",
         size=10, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)
    para(doc, "2026 年 10 月 23 日", size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.save(path)
    print("saved:", path)


if __name__ == "__main__":
    import sys
    build(sys.argv[1] if len(sys.argv) > 1 else "采购清单.docx")
